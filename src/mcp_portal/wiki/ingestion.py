"""Parse and ingest trusted local documents into persistent wiki storage."""

from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated
from urllib.parse import urlsplit

from bs4 import BeautifulSoup
from markdownify import markdownify

from mcp_portal.errors import ConfigurationPortalError, ValidationPortalError
from mcp_portal.wiki.models import (
    WikiCitation,
    WikiIngestionResult,
    WikiPage,
    WikiPageRecord,
    WikiPassageRecord,
    WikiSourceRecord,
)
from mcp_portal.wiki.repository import WikiRepository

SUPPORTED_SUFFIXES = frozenset({".docx", ".htm", ".html", ".md", ".markdown", ".pdf", ".txt"})
DEFAULT_MAX_BYTES = 25 * 1024 * 1024
MAX_PASSAGES = 2_000
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_LABEL = re.compile(r"^[a-z0-9]+(?:[._-][a-z0-9]+)*$")
_REMOVABLE_HTML = "script, style, noscript, template, svg, canvas, iframe, object, embed"


@dataclass(frozen=True)
class ParsedWikiDocument:
    """Normalized source document ready for deterministic chunking.

    Attributes:
        path: Resolved local source path.
        title: Parser-derived human-readable document title.
        markdown: Normalized Markdown representation.
        document_format: Stable parser format name.
        media_type: Source media type.
        source_updated_at: Last-modified time reported by the filesystem.
        byte_count: Number of source bytes read.
        source_content_hash: SHA-256 digest of the source bytes.
    """

    path: Annotated[Path, "Resolved local source path."]
    title: Annotated[str, "Parser-derived human-readable document title."]
    markdown: Annotated[str, "Normalized Markdown representation."]
    document_format: Annotated[str, "Stable parser format name."]
    media_type: Annotated[str, "Source media type."]
    source_updated_at: Annotated[datetime, "Last-modified time reported by the filesystem."]
    byte_count: Annotated[int, "Number of source bytes read."]
    source_content_hash: Annotated[str, "SHA-256 digest of the source bytes."]


@dataclass(frozen=True)
class WikiDocumentChunk:
    """One deterministic heading-aware retrieval passage.

    Attributes:
        index: Zero-based position in the normalized document.
        heading: Optional nearest Markdown heading.
        text: Bounded passage text.
    """

    index: Annotated[int, "Zero-based position in the normalized document."]
    heading: Annotated[str | None, "Optional nearest Markdown heading."]
    text: Annotated[str, "Bounded passage text."]


@dataclass(frozen=True)
class WikiIngestionOptions:
    """Trusted operator inputs controlling one local document ingestion.

    Attributes:
        path: Local document to parse.
        tenant_partition: Trusted non-reversible tenant partition.
        slug: Optional explicit wiki page slug.
        title: Optional explicit page and source title.
        source_id: Optional stable logical source identifier.
        source_uri: Optional canonical URI exposed by citations.
        tags: Tags attached to the page, source, and passages.
        required_scopes: Additional scopes required to retrieve the document.
        chunk_characters: Maximum characters in one retrieval passage.
        chunk_overlap: Characters repeated between adjacent passage windows.
        max_bytes: Maximum source size accepted by the parser.
        dry_run: Whether to validate without changing persistent state.
    """

    path: Annotated[Path, "Local document to parse."]
    tenant_partition: Annotated[str, "Trusted non-reversible tenant partition."]
    slug: Annotated[str | None, "Optional explicit wiki page slug."] = None
    title: Annotated[str | None, "Optional explicit page and source title."] = None
    source_id: Annotated[str | None, "Optional stable logical source identifier."] = None
    source_uri: Annotated[str | None, "Optional canonical URI exposed by citations."] = None
    tags: Annotated[tuple[str, ...], "Tags attached to the page, source, and passages."] = ()
    required_scopes: Annotated[
        frozenset[str], "Additional scopes required to retrieve the document."
    ] = frozenset()
    chunk_characters: Annotated[int, "Maximum characters in one retrieval passage."] = 1_800
    chunk_overlap: Annotated[int, "Characters repeated between adjacent passage windows."] = 200
    max_bytes: Annotated[int, "Maximum source size accepted by the parser."] = DEFAULT_MAX_BYTES
    dry_run: Annotated[bool, "Whether to validate without changing persistent state."] = False

    def __post_init__(self) -> None:
        """Validate bounded ingestion settings before reading a document."""
        if not self.tenant_partition:
            raise ValueError("A trusted tenant partition is required")
        if not 500 <= self.chunk_characters <= 8_000:
            raise ValueError("Chunk characters must be between 500 and 8000")
        if not 0 <= self.chunk_overlap < self.chunk_characters // 2:
            raise ValueError("Chunk overlap must be non-negative and less than half the chunk size")
        if not 1 <= self.max_bytes <= 100 * 1024 * 1024:
            raise ValueError("Maximum document size must be between 1 byte and 100 MiB")


def ingest_file(
    repository: WikiRepository | None,
    options: WikiIngestionOptions,
    *,
    ingested_at: datetime | None = None,
) -> WikiIngestionResult:
    """Parse, chunk, and atomically publish one trusted local document.

    Args:
        repository: Persistent repository, required unless this is a dry run.
        options: Trusted operator inputs and parsing limits.
        ingested_at: Optional deterministic ingestion timestamp.

    Returns:
        Sanitized identifiers, hashes, and counts for the ingestion.
    """
    parsed = parse_document(options.path, max_bytes=options.max_bytes)
    title = _normalize_title(options.title or parsed.title)
    slug = _normalize_slug(options.slug or parsed.path.stem)
    tags = _normalize_tags(options.tags)
    required_scopes = frozenset(_normalize_scopes(tuple(options.required_scopes)))
    source_id = _source_id(options.source_id, parsed.path)
    source_uri = _source_uri(options.source_uri, source_id)
    page_markdown = _page_markdown(title, parsed.markdown)
    page_content_hash = _content_hash(page_markdown.encode("utf-8"))
    page_revision_id = _page_revision_id(
        {
            "markdown": page_markdown,
            "required_scopes": sorted(required_scopes),
            "source_content_hash": parsed.source_content_hash,
            "source_uri": source_uri,
            "tags": tags,
            "title": title,
        }
    )
    chunks = chunk_markdown(
        page_markdown,
        chunk_characters=options.chunk_characters,
        overlap=options.chunk_overlap,
    )
    committed_at = ingested_at or datetime.now(timezone.utc)
    citations = tuple(
        _citation(parsed, chunk, source_id=source_id, source_uri=source_uri, title=title)
        for chunk in chunks
    )
    page = WikiPageRecord(
        tenant_partition=options.tenant_partition,
        required_scopes=required_scopes,
        page=WikiPage(
            slug=slug,
            revision_id=page_revision_id,
            title=title,
            summary=_summary(chunks),
            markdown=page_markdown,
            status="published",
            updated_at=committed_at,
            source_updated_at=parsed.source_updated_at,
            content_hash=page_content_hash,
            tags=list(tags),
            citations=list(citations),
        ),
    )
    source = WikiSourceRecord(
        tenant_partition=options.tenant_partition,
        source_id=source_id,
        source_revision=parsed.source_content_hash,
        source_uri=source_uri,
        title=title,
        document_format=parsed.document_format,
        content_hash=parsed.source_content_hash,
        source_updated_at=parsed.source_updated_at,
        ingested_at=committed_at,
        byte_count=parsed.byte_count,
        page_slug=slug,
        page_revision_id=page_revision_id,
        tags=tags,
        required_scopes=required_scopes,
    )
    passages = tuple(
        WikiPassageRecord(
            tenant_partition=options.tenant_partition,
            passage_id=_passage_id(source_id, parsed.source_content_hash, chunk.index),
            page_slug=slug,
            text=chunk.text,
            citation=citations[chunk.index],
            tags=tags,
            required_scopes=required_scopes,
        )
        for chunk in chunks
    )
    if not options.dry_run and repository is None:
        raise ConfigurationPortalError(
            "Persistent wiki repository is required to publish an ingestion."
        )
    if not options.dry_run and repository is not None:
        repository.ingest_source(source, page, passages)
    return WikiIngestionResult(
        source_id=source_id,
        source_revision=parsed.source_content_hash,
        source_uri=source_uri,
        page_slug=slug,
        page_revision_id=page_revision_id,
        source_content_hash=parsed.source_content_hash,
        page_content_hash=page_content_hash,
        document_format=parsed.document_format,
        byte_count=parsed.byte_count,
        passage_count=len(passages),
        dry_run=options.dry_run,
    )


def parse_document(path: Path, *, max_bytes: int = DEFAULT_MAX_BYTES) -> ParsedWikiDocument:
    """Read and normalize one supported local document.

    Args:
        path: Local Markdown, text, HTML, PDF, or DOCX file.
        max_bytes: Maximum source bytes accepted before parsing.

    Returns:
        Normalized Markdown and source metadata.
    """
    try:
        selected = path.expanduser().resolve()
        if not selected.is_file():
            raise ValidationPortalError("Wiki ingestion path must be a readable file.")
        stat = selected.stat()
    except ValidationPortalError:
        raise
    except OSError as error:
        raise ValidationPortalError(
            "Wiki ingestion path could not be inspected.",
            details={"error_type": type(error).__name__},
        ) from error
    suffix = selected.suffix.casefold()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValidationPortalError(
            "Wiki document format is not supported.",
            details={"suffix": suffix, "supported_suffixes": sorted(SUPPORTED_SUFFIXES)},
        )
    if stat.st_size <= 0 or stat.st_size > max_bytes:
        raise ValidationPortalError(
            "Wiki document size is outside the configured ingestion limit.",
            details={"byte_count": stat.st_size, "max_bytes": max_bytes},
        )
    try:
        raw = selected.read_bytes()
    except OSError as error:
        raise ValidationPortalError(
            "Wiki ingestion file could not be read.",
            details={"error_type": type(error).__name__},
        ) from error
    if not raw or len(raw) > max_bytes:
        raise ValidationPortalError(
            "Wiki document size changed outside the configured ingestion limit.",
            details={"byte_count": len(raw), "max_bytes": max_bytes},
        )
    parser = {
        ".docx": _parse_docx,
        ".htm": _parse_html,
        ".html": _parse_html,
        ".md": _parse_text,
        ".markdown": _parse_text,
        ".pdf": _parse_pdf,
        ".txt": _parse_text,
    }[suffix]
    document_format, media_type, parsed_title, markdown = parser(raw, selected)
    normalized = _normalize_markdown(markdown)
    return ParsedWikiDocument(
        path=selected,
        title=_normalize_title(parsed_title or selected.stem),
        markdown=normalized,
        document_format=document_format,
        media_type=media_type,
        source_updated_at=datetime.fromtimestamp(stat.st_mtime, timezone.utc),
        byte_count=len(raw),
        source_content_hash=_content_hash(raw),
    )


def chunk_markdown(
    markdown: str, *, chunk_characters: int = 1_800, overlap: int = 200
) -> tuple[WikiDocumentChunk, ...]:
    """Split Markdown into deterministic heading-aware retrieval passages.

    Args:
        markdown: Normalized page Markdown.
        chunk_characters: Maximum characters per passage.
        overlap: Characters repeated between adjacent windows in a section.

    Returns:
        Ordered non-empty retrieval passages.
    """
    sections: list[tuple[str | None, str]] = []
    heading: str | None = None
    body: list[str] = []
    for line in markdown.splitlines():
        match = _HEADING.fullmatch(line)
        if match is not None:
            if body:
                sections.append((heading, "\n".join(body).strip()))
            heading = match.group(2).strip()
            body = []
        else:
            body.append(line)
    if body:
        sections.append((heading, "\n".join(body).strip()))
    chunks: list[WikiDocumentChunk] = []
    for selected_heading, text in sections:
        for passage in _split_text(text, chunk_characters, overlap):
            chunks.append(WikiDocumentChunk(len(chunks), selected_heading, passage))
            if len(chunks) > MAX_PASSAGES:
                raise ValidationPortalError(
                    "Wiki document produced too many passages.",
                    details={"max_passages": MAX_PASSAGES},
                )
    if not chunks:
        raise ValidationPortalError("Wiki document did not contain extractable text.")
    return tuple(chunks)


def _parse_text(raw: bytes, path: Path) -> tuple[str, str, str | None, str]:
    """Decode a UTF-8 Markdown or plain-text document.

    Args:
        raw: Source document bytes.
        path: Resolved source path used to select format metadata.

    Returns:
        Format, media type, optional title, and Markdown content.
    """
    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValidationPortalError("Text wiki documents must use UTF-8 encoding.") from error
    is_markdown = path.suffix.casefold() in {".md", ".markdown"}
    return (
        "markdown" if is_markdown else "text",
        "text/markdown" if is_markdown else "text/plain",
        _first_heading(content),
        content,
    )


def _parse_html(raw: bytes, path: Path) -> tuple[str, str, str | None, str]:
    """Convert a local HTML document into normalized Markdown.

    Args:
        raw: Source HTML bytes.
        path: Resolved source path retained for the common parser contract.

    Returns:
        Format, media type, optional title, and Markdown content.
    """
    del path
    try:
        html = raw.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValidationPortalError("HTML wiki documents must use UTF-8 encoding.") from error
    soup = BeautifulSoup(html, "html.parser")
    title = " ".join(soup.title.stripped_strings) if soup.title is not None else None
    root = soup.find("main") or soup.find("article") or soup.body or soup
    for element in root.select(_REMOVABLE_HTML):
        element.decompose()
    content = markdownify(str(root), heading_style="ATX", bullets="-")
    return "html", "text/html", title or None, content


def _parse_pdf(raw: bytes, path: Path) -> tuple[str, str, str | None, str]:
    """Extract page-aware Markdown from a PDF document.

    Args:
        raw: Source PDF bytes.
        path: Resolved source path retained for the common parser contract.

    Returns:
        Format, media type, optional title, and Markdown content.
    """
    del path
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ConfigurationPortalError("PDF ingestion requires the pypdf dependency.") from error
    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValidationPortalError("Encrypted PDF documents are not supported.")
        pages = [
            f"## Page {index}\n\n{page.extract_text() or ''}"
            for index, page in enumerate(reader.pages, 1)
        ]
        title = (
            str(reader.metadata.title).strip()
            if reader.metadata and reader.metadata.title
            else None
        )
    except ValidationPortalError:
        raise
    except Exception as error:
        raise ValidationPortalError("PDF document could not be parsed.") from error
    return "pdf", "application/pdf", title, "\n\n".join(pages)


def _parse_docx(raw: bytes, path: Path) -> tuple[str, str, str | None, str]:
    """Convert Word paragraphs, headings, and tables into Markdown.

    Args:
        raw: Source DOCX bytes.
        path: Resolved source path retained for the common parser contract.

    Returns:
        Format, media type, optional title, and Markdown content.
    """
    del path
    try:
        from docx import Document
    except ImportError as error:
        raise ConfigurationPortalError(
            "DOCX ingestion requires the python-docx dependency."
        ) from error
    try:
        document = Document(io.BytesIO(raw))
        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            match = re.fullmatch(r"Heading (\d+)", style)
            lines.append(f"{'#' * min(int(match.group(1)), 6)} {text}" if match else text)
        for table in document.tables:
            rows = [[_table_cell(cell.text) for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            lines.extend(_markdown_table(rows))
        title = document.core_properties.title.strip() or None
    except Exception as error:
        raise ValidationPortalError("DOCX document could not be parsed.") from error
    return (
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        title,
        "\n\n".join(lines),
    )


def _markdown_table(rows: list[list[str]]) -> list[str]:
    """Render extracted Word table rows as Markdown.

    Args:
        rows: Table cells in source row order.

    Returns:
        Markdown table lines with a generated header separator.
    """
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    output = ["| " + " | ".join(normalized[0]) + " |"]
    output.append("| " + " | ".join("---" for _ in range(width)) + " |")
    output.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return output


def _split_text(text: str, limit: int, overlap: int) -> tuple[str, ...]:
    """Split one Markdown section near whitespace boundaries.

    Args:
        text: Section body without its heading line.
        limit: Maximum characters per passage.
        overlap: Characters repeated between adjacent windows.

    Returns:
        Ordered non-empty passage windows.
    """
    selected = text.strip()
    if not selected:
        return ()
    passages: list[str] = []
    cursor = 0
    while cursor < len(selected):
        end = min(len(selected), cursor + limit)
        if end < len(selected):
            boundary = max(
                selected.rfind("\n", cursor + limit // 2, end),
                selected.rfind(" ", cursor + limit // 2, end),
            )
            if boundary > cursor:
                end = boundary
        passage = selected[cursor:end].strip()
        if passage:
            passages.append(passage)
        if end >= len(selected):
            break
        cursor = max(cursor + 1, end - overlap)
    return tuple(passages)


def _normalize_markdown(value: str) -> str:
    """Normalize line endings and reject empty extracted content.

    Args:
        value: Parser-produced Markdown.

    Returns:
        Compact normalized Markdown.
    """
    selected = value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    selected = "\n".join(line.rstrip() for line in selected.splitlines()).strip()
    selected = re.sub(r"\n{3,}", "\n\n", selected)
    if not selected:
        raise ValidationPortalError("Wiki document did not contain extractable text.")
    return selected


def _normalize_title(value: str) -> str:
    """Normalize and validate a human-readable source title.

    Args:
        value: Parser- or operator-provided title.

    Returns:
        Single-line title bounded for persistent storage.
    """
    selected = " ".join(value.split())
    if not selected or len(selected) > 500:
        raise ValidationPortalError("Wiki document title must contain 1 to 500 characters.")
    return selected


def _normalize_slug(value: str) -> str:
    """Convert an operator value or filename into a valid page slug.

    Args:
        value: Requested slug or source filename stem.

    Returns:
        Lowercase URL-safe wiki page slug.
    """
    selected = re.sub(r"[^a-z0-9._-]+", "-", value.strip().casefold()).strip("._-")
    selected = re.sub(r"[._-]{2,}", "-", selected)
    if not selected or len(selected) > 200 or not _LABEL.fullmatch(selected):
        raise ValidationPortalError("Wiki page slug must contain 1 to 200 slug-like characters.")
    return selected


def _normalize_tags(values: tuple[str, ...]) -> tuple[str, ...]:
    """Normalize and validate document tags.

    Args:
        values: Operator-provided tags.

    Returns:
        Deduplicated normalized tags in input order.
    """
    selected = tuple(dict.fromkeys(value.strip().casefold() for value in values))
    if len(selected) > 50 or any(
        not value or len(value) > 160 or not _LABEL.fullmatch(value) for value in selected
    ):
        raise ValidationPortalError("Wiki tags must be slug-like values.")
    return selected


def _normalize_scopes(values: tuple[str, ...]) -> tuple[str, ...]:
    """Validate document scopes without changing identity-provider values.

    Args:
        values: Operator-provided authorization scope strings.

    Returns:
        Deduplicated scope strings in input order.
    """
    selected = tuple(dict.fromkeys(value.strip() for value in values))
    if len(selected) > 50 or any(
        not value or len(value) > 256 or any(character.isspace() for character in value)
        for value in selected
    ):
        raise ValidationPortalError(
            "Wiki required scopes must be non-empty values without whitespace."
        )
    return selected


def _source_id(value: str | None, path: Path) -> str:
    """Return an explicit or non-reversible path-derived source identifier.

    Args:
        value: Optional stable source identifier supplied by the operator.
        path: Resolved local source path.

    Returns:
        Stable source identifier safe for persistent metadata.
    """
    if value is not None:
        selected = value.strip()
        if not selected or len(selected) > 512:
            raise ValidationPortalError("Wiki source ID must contain 1 to 512 characters.")
        return selected
    digest = hashlib.sha256(str(path).casefold().encode("utf-8")).hexdigest()
    return f"local:{digest[:32]}"


def _source_uri(value: str | None, source_id: str) -> str:
    """Return a canonical public citation URI without leaking local paths.

    Args:
        value: Optional canonical URI supplied by the operator.
        source_id: Stable logical source identifier.

    Returns:
        Absolute citation URI or a generated non-path URN.
    """
    if value is None:
        digest = hashlib.sha256(source_id.encode("utf-8")).hexdigest()
        return f"urn:mcp-portal:wiki:source:{digest[:32]}"
    selected = value.strip()
    if len(selected) > 2_048 or not urlsplit(selected).scheme:
        raise ValidationPortalError("Wiki source URI must be an absolute URI.")
    return selected


def _page_markdown(title: str, markdown: str) -> str:
    """Ensure a source-backed page has one leading title heading.

    Args:
        title: Normalized page title.
        markdown: Normalized parser output.

    Returns:
        Publishable source-backed page Markdown.
    """
    first = markdown.splitlines()[0]
    return markdown if _HEADING.fullmatch(first) else f"# {title}\n\n{markdown}"


def _first_heading(markdown: str) -> str | None:
    """Return the first Markdown heading text when present.

    Args:
        markdown: Raw Markdown source text.

    Returns:
        First heading text, or None when no heading exists.
    """
    return next(
        (
            match.group(2).strip()
            for line in markdown.splitlines()
            if (match := _HEADING.fullmatch(line))
        ),
        None,
    )


def _summary(chunks: tuple[WikiDocumentChunk, ...]) -> str:
    """Create a bounded plain-text summary from the first passage.

    Args:
        chunks: Parsed document passages.

    Returns:
        Compact first-passage summary.
    """
    selected = re.sub(r"\s+", " ", chunks[0].text).strip()
    return selected[:2_000]


def _citation(
    document: ParsedWikiDocument,
    chunk: WikiDocumentChunk,
    *,
    source_id: str,
    source_uri: str,
    title: str,
) -> WikiCitation:
    """Create one source-revision citation for a retrieval passage.

    Args:
        document: Parsed source metadata and content hash.
        chunk: Passage receiving the evidence reference.
        source_id: Stable logical source identifier.
        source_uri: Canonical public citation URI.
        title: Normalized source title.

    Returns:
        Immutable citation bound to the source revision.
    """
    source_token = hashlib.sha256(source_id.encode("utf-8")).hexdigest()[:20]
    return WikiCitation(
        citation_id=f"cite:{source_token}:{document.source_content_hash[7:23]}:{chunk.index:05d}",
        source_id=source_id,
        source_revision=document.source_content_hash,
        title=title,
        heading=chunk.heading,
        source_uri=source_uri,
        source_updated_at=document.source_updated_at,
        content_hash=document.source_content_hash,
    )


def _passage_id(source_id: str, source_revision: str, index: int) -> str:
    """Create a deterministic passage identifier for one source revision.

    Args:
        source_id: Stable logical source identifier.
        source_revision: Content-derived source revision identifier.
        index: Passage position within the parsed document.

    Returns:
        Stable bounded passage identifier.
    """
    source_token = hashlib.sha256(source_id.encode("utf-8")).hexdigest()[:20]
    return f"passage:{source_token}:{source_revision[7:23]}:{index:05d}"


def _content_hash(value: bytes) -> str:
    """Return the canonical SHA-256 identifier for content bytes.

    Args:
        value: Source or normalized page bytes.

    Returns:
        SHA-256 identifier with the required prefix.
    """
    return f"sha256:{hashlib.sha256(value).hexdigest()}"


def _page_revision_id(envelope: dict[str, object]) -> str:
    """Hash the complete source-backed page revision envelope.

    Args:
        envelope: Canonical page content, provenance, tags, and authorization metadata.

    Returns:
        Content-addressed immutable page revision identifier.
    """
    serialized = json.dumps(
        envelope,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    )
    return _content_hash(serialized.encode("utf-8"))


def _table_cell(value: str) -> str:
    """Normalize one Word table cell for Markdown output.

    Args:
        value: Extracted Word cell text.

    Returns:
        Single-line Markdown-safe cell text.
    """
    return " ".join(value.split()).replace("|", "\\|")
